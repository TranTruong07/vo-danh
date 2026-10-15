import importlib.util
import tempfile
import unittest
from pathlib import Path
from docx import Document

spec = importlib.util.spec_from_file_location("extract_knowledge", "scripts/extract-knowledge.py")
extractor = importlib.util.module_from_spec(spec)
spec.loader.exec_module(extractor)


class ExtractionTests(unittest.TestCase):
    def test_body_table_order_and_locators(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fixture.docx"
            document = Document()
            document.add_paragraph("I. Mục thử nghiệm")
            document.add_paragraph("Đoạn trước bảng")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Ô A"
            table.cell(0, 1).text = "Ô B"
            table.cell(1, 0).merge(table.cell(1, 1)).text = "Ô gộp"
            document.add_paragraph("Đoạn sau bảng")
            document.add_paragraph("IX. Các điểm cần phân biệt")
            document.add_paragraph("1. Nội dung cần giữ nguyên.")
            document.save(path)
            result = extractor.extract(path)
            self.assertEqual([b["text"] for b in result["blocks"]], ["I. Mục thử nghiệm", "Đoạn trước bảng", "Ô A", "Ô B", "Ô gộp", "Đoạn sau bảng", "IX. Các điểm cần phân biệt", "1. Nội dung cần giữ nguyên."])
            self.assertEqual(result["counts"]["tables"], 1)
            self.assertEqual(result["blocks"][2]["tableIndex"], 0)
            self.assertEqual(result["blocks"][2]["rowIndex"], 0)
            self.assertEqual(result["blocks"][-1]["sectionId"], "IX.1")
            self.assertFalse(result["blocks"][-1]["isHeading"])
            self.assertEqual(result, extractor.extract(path))

    def test_actual_source_keeps_all_sections_table_and_correction(self):
        data = extractor.extract("docs/source/Tong-hop-kien-thuc.docx")
        self.assertEqual([section["sectionId"] for section in data["sections"]], ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"])
        self.assertEqual(data["counts"]["paragraphs"], 189)
        self.assertEqual(data["counts"]["tables"], 1)
        correction = next(block for block in data["blocks"] if block["text"].startswith("Đính chính quan trọng"))
        last_cell = [block for block in data["blocks"] if "tableIndex" in block][-1]
        self.assertGreater(correction["order"], last_cell["order"])
        self.assertIn("Côn Đảo", correction["text"])
        self.assertTrue(any(block["sectionId"] == "IX.4" for block in data["blocks"]))


if __name__ == "__main__":
    unittest.main()
